import openpyxl
import pytest

from chunknorris.core.components import MarkdownDoc
from chunknorris.parsers import ExcelParser


def test_read_file(excel_parser: ExcelParser, excel_filepath: str):
    sheets = excel_parser.read_file(excel_filepath)
    assert len(sheets) == 2
    df = sheets["CSV-like"]
    assert len(df.index) == 5 and len(df.columns) == 6  # type: ignore | missing typing in pd.Index


def test_parse_file(excel_parser: ExcelParser, excel_filepath: str):
    # output format = auto -> 1 sheet should be parsed as JSON lines, the other as MD table
    excel_parser.output_format = "auto"
    parser_output = excel_parser.parse_file(excel_filepath)
    assert isinstance(parser_output, MarkdownDoc)
    lines = parser_output.to_string().split("\n")
    assert sum(line.startswith("#") for line in lines) == 2
    assert sum(line.startswith("|") for line in lines) == 7
    assert sum(line.startswith("{") for line in lines) == 5


def test_parse_string(excel_parser: ExcelParser, excel_filepath: str):
    excel_parser.output_format = "auto"
    with open(excel_filepath, "rb") as f:
        file_content = f.read()
    parser_output = excel_parser.parse_string(file_content)
    assert isinstance(parser_output, MarkdownDoc)
    lines = parser_output.to_string().split("\n")
    assert sum(line.startswith("#") for line in lines) == 2
    assert sum(line.startswith("|") for line in lines) == 7
    assert sum(line.startswith("{") for line in lines) == 5


def test_inspect_file(excel_parser: ExcelParser, excel_filepath: str):
    profiles = excel_parser.inspect_file(excel_filepath)
    assert list(profiles) == ["CSV-like", "Not CSV-like"]
    assert profiles["CSV-like"].is_csv_like
    assert not profiles["Not CSV-like"].is_csv_like
    assert profiles["CSV-like"].n_rows == 5


def test_inspect_string(excel_parser: ExcelParser, excel_filepath: str):
    with open(excel_filepath, "rb") as f:
        file_content = f.read()
    assert excel_parser.inspect_string(file_content) == excel_parser.inspect_file(
        excel_filepath
    )


def test_inspect_file_samples_rows(tmp_path):
    filepath = tmp_path / "file.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["id", "name"])
    for i in range(3000):
        worksheet.append([i, f"name {i}"])
    workbook.save(filepath)
    profile = ExcelParser(n_sample_rows=10).inspect_file(str(filepath))["Sheet"]
    assert profile.n_rows == 3000
    assert profile.is_csv_like


def test_inspect_file_without_dimension_metadata(tmp_path):
    filepath = tmp_path / "file.xlsx"
    workbook = openpyxl.Workbook(write_only=True)  # does not write <dimension>
    worksheet = workbook.create_sheet("Sheet")
    worksheet.append(["id", "name"])
    for i in range(20):
        worksheet.append([i, f"name {i}"])
    workbook.save(filepath)
    profile = ExcelParser().inspect_file(str(filepath))["Sheet"]
    assert profile.n_rows is None
    assert profile.is_csv_like


def test_inspect_file_wrong_extension(excel_parser: ExcelParser):
    with pytest.raises(ValueError):
        excel_parser.inspect_file("file.csv")


def test_parse_file_stores_sheet_profiles(excel_filepath: str):
    parser = ExcelParser(output_format="markdown_table")
    parser.parse_file(excel_filepath)
    assert parser.sheet_profiles == parser.inspect_file(excel_filepath)
    assert parser.sheet_profiles["CSV-like"].n_rows == 5
