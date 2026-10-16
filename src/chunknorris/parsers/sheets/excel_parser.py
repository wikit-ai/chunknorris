from io import BytesIO
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile

import openpyxl
import pandas as pd
from openpyxl.utils.exceptions import InvalidFileException

from ...core.components import MarkdownDoc
from .abstract_sheet_parser import AbstractSheetParser
from .sheet_profile import SheetProfile

_VALID_EXTENSIONS = (".xls", ".xlsx", ".xlsm", ".xlsb", ".odf", ".ods", ".odt")


class ExcelParser(AbstractSheetParser[bytes]):
    """Parser for spreadsheets, such as Excel workbooks (.xslx). For a list of handled filetypes,
    refer to https://pandas.pydata.org/docs/reference/api/pandas.read_excel.html"""

    def __init__(
        self,
        output_format: Literal["markdown_table", "json_lines", "auto"] = "auto",
        n_sample_rows: int = 1000,
    ) -> None:
        """Initializes an Excel parser

        Args:
            output_format (Literal["markdown_table", "json_lines", "auto], optional): the output format of the parsed document.
                - markdown_table : uses tabula to build a markdown-formatted table.
                - json_lines : each row of the table will be output as a JSON line. Better for chunking as headers are preserved.
                - auto : will detect which format is the more suitable. CSV-like sheet will be converset to JSON lines.
                Defaults to "auto".
            n_sample_rows (int, optional): number of rows sampled at the top of each sheet
                to compute its SheetProfile. Defaults to 1000.
        """
        self.output_format = output_format
        self.n_sample_rows = n_sample_rows
        # Profiles of the sheets of the last parsed file. See SheetProfile.
        self.sheet_profiles: dict[str, SheetProfile] = {}

    def parse_file(self, filepath: str) -> MarkdownDoc:
        """Parses a excel-like file to markdown.

        Args:
            filepath (str): the path to the excel-like file.

        Returns:
            MarkdownDoc: the markdown formatted excel file.
        """
        sheets = self.read_file(filepath)
        md_string = self.convert_sheets_to_output_format(sheets)

        return MarkdownDoc.from_string(md_string)

    def parse_string(self, string: bytes) -> MarkdownDoc:
        """Parses a bytes string representing an excel file.

        Args:
            string (bytes): the excel as a byte string.

        Returns:
            MarkdownDoc: the markdown formatted excel file
        """
        sheets = pd.read_excel(BytesIO(string), sheet_name=None)  # type: ignore | missing typing in pandas.
        md_string = self.convert_sheets_to_output_format(sheets)

        return MarkdownDoc.from_string(md_string)

    def read_file(self, filepath: str) -> dict[str, pd.DataFrame]:
        """Read the provided filepath.

        Args:
            filepath (str): path to the file.

        Returns:
            dict[str, pd.DataFrame]: a mapping containing {sheet_name: corresponding_dataframe.}
        """
        ExcelParser._check_extension(filepath)
        sheets = pd.read_excel(filepath, sheet_name=None)  # type: ignore | missing typing in pandas.

        return sheets

    def inspect_file(self, filepath: str) -> dict[str, SheetProfile]:
        """Describes the structure of each sheet of an excel-like file, without loading it entirely.
        Useful to route the file before parsing, e.g. to handle long CSV-like sheets differently.

        Args:
            filepath (str): path to the file.

        Returns:
            dict[str, SheetProfile]: a mapping containing {sheet_name: corresponding_profile}.
        """
        ExcelParser._check_extension(filepath)
        return self._inspect(filepath)

    def inspect_string(self, string: bytes) -> dict[str, SheetProfile]:
        """Describes the structure of each sheet of an excel file, without loading it entirely.
        Useful to route the file before parsing, e.g. to handle long CSV-like sheets differently.

        Args:
            string (bytes): the excel as a byte string.

        Returns:
            dict[str, SheetProfile]: a mapping containing {sheet_name: corresponding_profile}.
        """
        return self._inspect(BytesIO(string))

    def _inspect(self, source: str | BytesIO) -> dict[str, SheetProfile]:
        n_rows: dict[str, int | None] = {}
        # For the formats it handles (.xlsx, .xlsm), the workbook is opened once with openpyxl
        # to read the row counts, then handed to pandas, which only reads the sampled rows.
        # Same options as pandas uses.
        try:
            workbook = openpyxl.load_workbook(
                source, read_only=True, data_only=True, keep_links=False
            )
        except (InvalidFileException, BadZipFile, KeyError):
            if isinstance(source, BytesIO):
                source.seek(0)
            samples = pd.read_excel(source, sheet_name=None, nrows=self.n_sample_rows)  # type: ignore | missing typing in pandas.
        else:
            try:
                # max_row is read from the <dimension> tag, which Excel always writes.
                n_rows = {
                    worksheet.title: (
                        max(worksheet.max_row - 1, 0)
                        if worksheet.max_row is not None
                        else None
                    )
                    for worksheet in workbook.worksheets
                }
                samples = pd.read_excel(workbook, sheet_name=None, nrows=self.n_sample_rows, engine="openpyxl")  # type: ignore | missing typing in pandas.
            finally:
                workbook.close()

        profiles: dict[str, SheetProfile] = {}
        for sheet_name, df in samples.items():
            sheet_n_rows = n_rows.get(str(sheet_name))
            # The metadata is wrong if the sample holds more rows than announced.
            if sheet_n_rows is not None and sheet_n_rows < len(df):
                sheet_n_rows = None
            profiles[str(sheet_name)] = SheetProfile.from_dataframe(
                str(sheet_name), df, sheet_n_rows
            )

        return profiles

    @staticmethod
    def _check_extension(filepath: str) -> None:
        suffix = Path(filepath).suffix.lower()
        if suffix not in _VALID_EXTENSIONS:
            raise ValueError(f"ExcelParser cannot parse {suffix} files.")

    def convert_sheets_to_output_format(self, sheets: dict[str, pd.DataFrame]) -> str:
        """Handle the conversion of the sheets obtained from
        pandas.read_excel() method to the specified output format.
        The profile of each sheet is stored in self.sheet_profiles.

        Args:
            sheets (dict[str, pd.DataFrame]): the sheets returned from pd.read_excel(sheet_name=None).

        Returns:
            str: the formatted string
        """
        output_string = ""
        self.sheet_profiles = {}

        for sheet_name, df in sheets.items():
            output_string += f"## {sheet_name}\n\n"

            profile = SheetProfile.from_dataframe(
                str(sheet_name), df.head(self.n_sample_rows), n_rows=len(df)
            )
            self.sheet_profiles[str(sheet_name)] = profile

            format_to_use = self.output_format
            if self.output_format == "auto":
                # CSV-like sheets are better chunked as JSON lines, as headers are repeated.
                format_to_use = (
                    "json_lines" if profile.is_csv_like else "markdown_table"
                )

            match format_to_use:
                case "markdown_table":
                    output_string += self.convert_df_to_markdown_table(df) + "\n\n"
                case "json_lines":
                    output_string += self.convert_df_to_json_lines(df) + "\n\n"
                case _:
                    raise ValueError(
                        f"Invalid value for argument 'output_format': expected one of ['markdown_table', 'json_lines']. Got '{format_to_use}'."
                    )

        return output_string
